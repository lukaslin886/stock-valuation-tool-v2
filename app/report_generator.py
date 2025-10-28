"""
報告生成模組
負責生成 Excel 和 PDF 格式的綜合分析報告
"""

import pandas as pd
import numpy as np
from datetime import datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image as RLImage
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import plotly.graph_objects as go
from PIL import Image as PILImage


class ReportGenerator:
    """報告生成器"""
    
    def __init__(self):
        """初始化報告生成器"""
        # 註冊中文字體（使用 Windows 系統字體）
        try:
            # 微軟正黑體
            font_path = "C:/Windows/Fonts/msjh.ttc"
            pdfmetrics.registerFont(TTFont('MSJH', font_path))
            self.chinese_font = 'MSJH'
            print(f"✓ 成功註冊中文字體: {self.chinese_font}")
        except Exception as e:
            print(f"⚠️ 字體註冊失敗，嘗試備用字體: {str(e)}")
            try:
                # 備用：標楷體
                font_path = "C:/Windows/Fonts/kaiu.ttf"
                pdfmetrics.registerFont(TTFont('Kaiu', font_path))
                self.chinese_font = 'Kaiu'
                print(f"✓ 成功註冊備用中文字體: {self.chinese_font}")
            except Exception as e2:
                print(f"❌ 所有字體註冊失敗: {str(e2)}")
                self.chinese_font = 'Helvetica'  # 最後使用預設字體
    
    def generate_excel_report(
        self,
        stock_code: str,
        stock_name: str,
        dcf_result: dict,
        risk_report: dict,
        current_price: float,
        current_eps: float,
        investment_amount: float = 100000
    ) -> BytesIO:
        """
        生成 Excel 格式報告
        
        Args:
            stock_code: 股票代碼
            stock_name: 股票名稱
            dcf_result: DCF 計算結果
            risk_report: 風險評估報告
            current_price: 目前股價
            current_eps: 目前 EPS
            investment_amount: 投資金額
            
        Returns:
            BytesIO: Excel 檔案內容
        """
        # 建立 Excel 工作簿
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        
        # 1. 建立摘要工作表
        self._create_summary_sheet(
            wb, stock_code, stock_name, dcf_result, risk_report,
            current_price, current_eps, investment_amount
        )
        
        # 2. 建立 DCF 詳細分析工作表
        self._create_dcf_detail_sheet(wb, dcf_result)
        
        # 3. 建立風險分析工作表
        if 'error' not in risk_report:
            self._create_risk_analysis_sheet(wb, risk_report)
        
        # 儲存到 BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    def _create_summary_sheet(
        self, wb, stock_code, stock_name, dcf_result, risk_report,
        current_price, current_eps, investment_amount
    ):
        """建立摘要工作表"""
        ws = wb.create_sheet("摘要", 0)
        
        # 設定欄寬
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 15
        
        # 標題
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f'{stock_code} {stock_name} 綜合分析報告'
        title_cell.font = Font(size=16, bold=True, color='FFFFFF')
        title_cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # 報告日期
        ws.merge_cells('A2:D2')
        date_cell = ws['A2']
        date_cell.value = f'報告日期：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
        date_cell.alignment = Alignment(horizontal='center')
        
        # 基本資訊
        row = 4
        self._add_section_header(ws, f'A{row}', 'D', '基本資訊')
        
        row += 1
        data = [
            ['股票代碼', stock_code, '目前股價', f'${current_price:.2f}'],
            ['股票名稱', stock_name, '最新 EPS', f'${current_eps:.2f}'],
            ['投資金額', f'${investment_amount:,.0f}', '本益比', f'{current_price/current_eps:.2f}' if current_eps > 0 else 'N/A']
        ]
        
        for row_data in data:
            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                if col % 2 == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
            row += 1
        
        # DCF 估值結果
        row += 1
        self._add_section_header(ws, f'A{row}', 'D', 'DCF 估值結果')
        
        row += 1
        dcf_data = [
            ['內在價值', f'${dcf_result["intrinsic_value"]:.2f}', '目前股價', f'${dcf_result["current_price"]:.2f}'],
            ['潛在獲利率', f'{dcf_result["upside_potential"]:.2%}', '折現率', f'{dcf_result["discount_rate"]:.2%}'],
            ['投資建議', dcf_result['recommendation'], '', '']
        ]
        
        for row_data in dcf_data:
            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                if col % 2 == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
                if col == 2 and row_data[0] == '潛在獲利率':
                    upside = dcf_result["upside_potential"]
                    if upside > 0.2:
                        cell.font = Font(bold=True, color='00B050')
                    elif upside < 0:
                        cell.font = Font(bold=True, color='FF0000')
            row += 1
        
        # 風險評估
        if 'error' not in risk_report and 'risk_score' in risk_report:
            row += 1
            self._add_section_header(ws, f'A{row}', 'D', '風險評估')
            
            row += 1
            risk_score = risk_report['risk_score']
            risk_data = [['風險評分', f'{risk_score["score"]:.1f}/100', '風險等級', risk_score['level']]]
            
            if 'volatility_analysis' in risk_report:
                vol = risk_report['volatility_analysis']
                risk_data.append(['年化波動率', f'{vol.get("annual_volatility", 0):.2%}', '', ''])
            
            if 'beta_analysis' in risk_report:
                beta = risk_report['beta_analysis']
                risk_data.append(['Beta 係數', f'{beta.get("beta", 0):.2f}', '', ''])
            
            for row_data in risk_data:
                for col, value in enumerate(row_data, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    if col % 2 == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
                row += 1
        
        # 添加邊框
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        for row_cells in ws.iter_rows(min_row=4, max_row=row-1, min_col=1, max_col=4):
            for cell in row_cells:
                cell.border = thin_border
    
    def _create_dcf_detail_sheet(self, wb, dcf_result):
        """建立 DCF 詳細分析工作表"""
        ws = wb.create_sheet("DCF詳細分析")
        
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18
        
        self._add_section_header(ws, 'A1', 'C', 'DCF 現金流預測')
        
        row = 2
        headers = ['年度', '預測現金流', '現值']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        row = 3
        for i, (cf, pv) in enumerate(zip(dcf_result['cash_flows'], dcf_result['present_values'])):
            ws.cell(row=row, column=1, value=f'第{i+1}年')
            ws.cell(row=row, column=2, value=cf).number_format = '#,##0.00'
            ws.cell(row=row, column=3, value=pv).number_format = '#,##0.00'
            row += 1
        
        ws.cell(row=row, column=1, value='總計').font = Font(bold=True)
        ws.cell(row=row, column=2, value=sum(dcf_result['cash_flows'])).number_format = '#,##0.00'
        ws.cell(row=row, column=3, value=sum(dcf_result['present_values'])).number_format = '#,##0.00'
        
        self._add_borders(ws, 2, row, 1, 3)
    
    def _create_risk_analysis_sheet(self, wb, risk_report):
        """建立風險分析工作表"""
        ws = wb.create_sheet("風險分析")
        
        row = 1
        
        if 'value_at_risk' in risk_report and 'error' not in risk_report['value_at_risk']:
            self._add_section_header(ws, f'A{row}', 'D', 'Value at Risk (VaR) 分析')
            row += 1
            
            var = risk_report['value_at_risk']
            headers = ['方法', 'VaR 金額', '百分比']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            
            row += 1
            var_methods = [
                ('參數法', var['parametric_var']),
                ('歷史模擬法', var['historical_var']),
                ('CVaR', var['cvar'])
            ]
            
            for method_name, method_data in var_methods:
                ws.cell(row=row, column=1, value=method_name)
                ws.cell(row=row, column=2, value=method_data['amount']).number_format = '$#,##0'
                ws.cell(row=row, column=3, value=method_data['percentage']).number_format = '0.00%'
                row += 1
        
        self._add_borders(ws, 2, row-1, 1, 3)
    
    def _add_section_header(self, ws, start_cell, end_col, title):
        """添加區塊標題"""
        ws.merge_cells(f'{start_cell}:{end_col}{start_cell[1:]}')
        cell = ws[start_cell]
        cell.value = title
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        cell.font = Font(bold=True, size=12, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    def _add_borders(self, ws, start_row, end_row, start_col, end_col):
        """添加邊框"""
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).border = thin_border
    
    def generate_pdf_report(
        self,
        stock_code: str,
        stock_name: str,
        dcf_result: dict,
        risk_report: dict,
        current_price: float,
        current_eps: float,
        investment_amount: float = 100000,
        chart_images: dict = None
    ) -> BytesIO:
        """
        生成 PDF 格式報告
        
        Args:
            stock_code: 股票代碼
            stock_name: 股票名稱
            dcf_result: DCF 計算結果
            risk_report: 風險評估報告
            current_price: 目前股價
            current_eps: 目前 EPS
            investment_amount: 投資金額
            chart_images: 圖表圖片字典 {name: BytesIO}
            
        Returns:
            BytesIO: PDF 檔案內容
        """
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # 自訂樣式（使用中文字體）
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontName=self.chinese_font,
            fontSize=24,
            textColor=colors.HexColor('#4472C4'),
            spaceAfter=30,
            alignment=TA_CENTER
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontName=self.chinese_font,
            fontSize=16,
            textColor=colors.HexColor('#4472C4'),
            spaceAfter=12,
            spaceBefore=12
        )
        
        # 自訂一般文字樣式
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontName=self.chinese_font,
            fontSize=11
        )
        
        # 標題
        title = Paragraph(f'{stock_code} {stock_name} 綜合分析報告', title_style)
        story.append(title)
        
        date_text = Paragraph(
            f'報告日期：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
            normal_style
        )
        story.append(date_text)
        story.append(Spacer(1, 0.3*inch))
        
        # 基本資訊
        story.append(Paragraph('基本資訊', heading_style))
        
        basic_data = [
            ['項目', '數值', '項目', '數值'],
            ['股票代碼', stock_code, '目前股價', f'${current_price:.2f}'],
            ['股票名稱', stock_name, '最新 EPS', f'${current_eps:.2f}'],
            ['投資金額', f'${investment_amount:,.0f}', '本益比', 
             f'{current_price/current_eps:.2f}' if current_eps > 0 else 'N/A']
        ]
        
        basic_table = Table(basic_data, colWidths=[1.5*inch, 1.5*inch, 1.5*inch, 1.5*inch])
        basic_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, -1), self.chinese_font),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('FONTSIZE', (0, 1), (-1, -1), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#E7E6E6')),
            ('BACKGROUND', (2, 1), (2, -1), colors.HexColor('#E7E6E6')),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(basic_table)
        story.append(Spacer(1, 0.3*inch))
        
        # DCF 估值結果
        story.append(Paragraph('DCF 估值結果', heading_style))
        
        dcf_data = [
            ['項目', '數值'],
            ['內在價值', f'${dcf_result["intrinsic_value"]:.2f}'],
            ['目前股價', f'${dcf_result["current_price"]:.2f}'],
            ['潛在獲利率', f'{dcf_result["upside_potential"]:.2%}'],
            ['折現率', f'{dcf_result["discount_rate"]:.2%}'],
            ['投資建議', dcf_result['recommendation']]
        ]
        
        dcf_table = Table(dcf_data, colWidths=[2*inch, 4*inch])
        dcf_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), self.chinese_font),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('FONTSIZE', (0, 1), (-1, -1), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#E7E6E6')),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(dcf_table)
        story.append(Spacer(1, 0.3*inch))
        
        # 插入圖表（如果有）
        if chart_images and 'cash_flow_chart' in chart_images:
            story.append(Paragraph('現金流預測圖', heading_style))
            try:
                img = RLImage(chart_images['cash_flow_chart'], width=5*inch, height=3*inch)
                story.append(img)
                story.append(Spacer(1, 0.2*inch))
            except:
                pass
        
        # 風險評估
        if 'error' not in risk_report and 'risk_score' in risk_report:
            story.append(PageBreak())
            story.append(Paragraph('風險評估', heading_style))
            
            risk_score = risk_report['risk_score']
            risk_data = [
                ['項目', '數值'],
                ['風險評分', f'{risk_score["score"]:.1f}/100'],
                ['風險等級', risk_score['level']]
            ]
            
            if 'volatility_analysis' in risk_report:
                vol = risk_report['volatility_analysis']
                risk_data.append(['年化波動率', f'{vol.get("annual_volatility", 0):.2%}'])
            
            if 'beta_analysis' in risk_report:
                beta = risk_report['beta_analysis']
                risk_data.append(['Beta 係數', f'{beta.get("beta", 0):.2f}'])
            
            risk_table = Table(risk_data, colWidths=[2*inch, 4*inch])
            risk_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, -1), self.chinese_font),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('FONTSIZE', (0, 1), (-1, -1), 11),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#E7E6E6')),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(risk_table)
        
        # 免責聲明
        story.append(Spacer(1, 0.5*inch))
        disclaimer = Paragraph(
            '<b>免責聲明：</b>本報告僅供參考，不構成投資建議。投資有風險，請謹慎評估。',
            normal_style
        )
        story.append(disclaimer)
        
        # 生成 PDF
        doc.build(story)
        buffer.seek(0)
        
        return buffer
    
    def save_plotly_chart_as_image(self, fig: go.Figure, width: int = 800, height: int = 600) -> BytesIO:
        """
        將 Plotly 圖表儲存為圖片
        
        Args:
            fig: Plotly 圖表物件
            width: 圖片寬度
            height: 圖片高度
            
        Returns:
            BytesIO: 圖片內容
        """
        try:
            img_bytes = fig.to_image(format="png", width=width, height=height)
            return BytesIO(img_bytes)
        except Exception as e:
            print(f"圖表轉換失敗: {str(e)}")
            return None
