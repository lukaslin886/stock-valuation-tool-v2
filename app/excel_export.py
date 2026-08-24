"""
Excel 匯出模組（帶進度回報）

把 DataFrame 分批寫入 .xlsx，並透過 progress_callback 回報進度，
供 Streamlit 端渲染「一鍵導出」進度條。與 UI 解耦，可獨立測試。

對應 TODO.md「增加『一鍵導出 Excel』進度條」。
"""

from io import BytesIO
from typing import Callable, Optional, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ProgressCB = Optional[Callable[[int, int], None]]


def build_screener_excel(
    df: pd.DataFrame,
    sheet_name: str = "篩選結果",
    batch_size: int = 50,
    progress_callback: ProgressCB = None,
    freeze_header: bool = True,
) -> bytes:
    """
    將 df 寫成 .xlsx bytes，分批回報進度。

    Args:
        df: 要匯出的資料（欄名即表頭）。
        sheet_name: 工作表名稱。
        batch_size: 每寫入幾列回報一次進度。
        progress_callback: fn(current_rows, total_rows)，current 會含表頭步驟。
        freeze_header: 是否凍結首列。

    Returns:
        .xlsx 檔的 bytes。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] if sheet_name else "Sheet1"

    cols: Sequence = list(df.columns)
    total = len(df) + 1  # +1 為表頭步驟
    done = 0

    def _tick(n_done: int):
        if progress_callback:
            progress_callback(min(n_done, total), total)

    # 1) 表頭
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for c_idx, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c_idx, value=str(col))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if freeze_header:
        ws.freeze_panes = "A2"
    done += 1
    _tick(done)

    # 2) 資料列（分批回報）
    values = df.itertuples(index=False, name=None)
    for r_idx, row in enumerate(values, start=2):
        for c_idx, val in enumerate(row, start=1):
            # NaN / numpy 型別轉為 openpyxl 可接受值
            if pd.isna(val):
                val = None
            elif hasattr(val, "item"):
                val = val.item()
            ws.cell(row=r_idx, column=c_idx, value=val)
        done += 1
        if (done % batch_size == 0) or (r_idx - 1 == len(df)):
            _tick(done)

    # 3) 欄寬自適應（依表頭與樣本估算）
    sample = df.head(200)
    for c_idx, col in enumerate(cols, start=1):
        max_len = len(str(col))
        if not sample.empty:
            col_max = sample.iloc[:, c_idx - 1].astype(str).map(len).max()
            max_len = max(max_len, int(col_max) if pd.notna(col_max) else 0)
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max(max_len + 2, 8), 40)

    _tick(total)  # 確保結束時 100%

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
