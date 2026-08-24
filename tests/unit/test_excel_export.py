"""
excel_export 單元測試：驗證 bytes 產出、內容正確、進度回報單調且收斂到 100%。
"""

import io
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "app"))

from excel_export import build_screener_excel  # noqa: E402


def _df():
    return pd.DataFrame({
        "代碼": ["2330", "2317", "2454"],
        "名稱": ["台積電", "鴻海", "聯發科"],
        "股價": [1000.5, 203.0, np.nan],
        "ROE": [0.28, 0.15, 0.22],
    })


def test_returns_valid_xlsx_bytes():
    data = build_screener_excel(_df())
    assert isinstance(data, (bytes, bytearray)) and len(data) > 0
    wb = load_workbook(io.BytesIO(data))
    assert wb.active.max_row == 4  # header + 3 rows


def test_header_and_values():
    wb = load_workbook(io.BytesIO(build_screener_excel(_df(), sheet_name="篩選結果")))
    ws = wb.active
    assert ws.title == "篩選結果"
    assert [ws.cell(1, c).value for c in range(1, 5)] == ["代碼", "名稱", "股價", "ROE"]
    assert ws.cell(2, 1).value == "2330"
    assert ws.freeze_panes == "A2"


def test_nan_becomes_empty():
    ws = load_workbook(io.BytesIO(build_screener_excel(_df()))).active
    assert ws.cell(4, 3).value in (None, "")  # NaN 股價


def test_progress_monotonic_and_complete():
    calls = []
    build_screener_excel(_df(), batch_size=1, progress_callback=lambda c, t: calls.append((c, t)))
    assert calls, "progress_callback should be invoked"
    assert all(calls[i][0] <= calls[i + 1][0] for i in range(len(calls) - 1))
    assert calls[0][0] == 1                      # 表頭先回報
    assert calls[-1][0] == calls[-1][1] == 4     # 收斂到 total = rows + header


def test_long_sheet_name_truncated():
    ws = load_workbook(io.BytesIO(build_screener_excel(_df(), sheet_name="x" * 40))).active
    assert len(ws.title) <= 31


def test_empty_dataframe():
    data = build_screener_excel(pd.DataFrame({"A": []}))
    assert isinstance(data, (bytes, bytearray)) and len(data) > 0
